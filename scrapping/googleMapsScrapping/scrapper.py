import os
import time
from fpdf import FPDF
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.action_chains import ActionChains
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import TimeoutException, NoSuchElementException
from webdriver_manager.chrome import ChromeDriverManager
from openpyxl import Workbook


# URL da página de avaliações que será raspada. Ficou como constante do
# módulo (em vez de uma variável solta dentro do main) porque é uma
# configuração do script, não um passo do processamento.
REVIEWS_URL = (
    'https://www.google.com/maps/place/Hospital+Erastinho/'
    '@-25.452814,-49.2411939,1074m/data=!3m1!1e3!4m8!3m7!1s0x94dce5124cb3e99b:'
    '0xdb43625e49c502d9!8m2!3d-25.4528189!4d-49.238619!9m1!1b1!16s%2Fg%2F11h5mp07v2'
    '?entry=ttu&g_ep=EgoyMDI2MDYyOC4wIKXMDSoASAFQAw%3D%3D'
)


# ---------------------------------------------------------------------------
# Utilitários gerais
# ---------------------------------------------------------------------------

def safe_text(text):
    """Garante que o texto seja codificado corretamente para uso no FPDF."""
    return text.encode('latin-1', 'replace').decode('latin-1')


def _real_click(driver, element):
    """
    Clica em um elemento simulando uma interação real de mouse (mousedown +
    mouseup + click), usando ActionChains.

    IMPORTANTE: isso resolve um problema comum com Selenium em páginas como o
    Google Maps. Um clique feito via
        driver.execute_script("arguments[0].click();", element)
    dispara APENAS o evento "click" em JavaScript. Só que o Google Maps abre
    seus menus a partir de eventos de "mousedown"/"pointerdown", não do
    "click" puro. Resultado: o Selenium não dá nenhum erro, mas o menu nunca
    chega a abrir de verdade, porque o evento certo nunca foi disparado.

    ActionChains simula o mouse "de verdade" (move até o elemento, pressiona
    o botão, solta o botão), então dispara toda a sequência de eventos que o
    Google Maps espera.
    """
    driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", element)
    time.sleep(0.3)
    ActionChains(driver).move_to_element(element).pause(0.2).click().perform()


# ---------------------------------------------------------------------------
# Configuração e carregamento inicial da página
# ---------------------------------------------------------------------------

def configure_driver(headless=False):
    """Cria e retorna uma instância configurada do Chrome WebDriver."""
    options = Options()
    if headless:
        options.add_argument('--headless')
    return webdriver.Chrome(service=Service(ChromeDriverManager().install()), options=options)


def accept_cookie_consent(driver, debug=True):
    """
    Antes de exibir o Google Maps normalmente, o Google às vezes mostra uma
    tela de consentimento de cookies ("Antes de continuar..." / "Aceitar
    tudo"). Enquanto essa tela não é fechada, a página do mapa/avaliações
    nunca termina de carregar de verdade.

    Essa função tenta localizar e clicar no botão de aceite. Se não
    encontrar nada, assume que a tela não apareceu dessa vez (o Google nem
    sempre mostra isso) e segue em frente normalmente.
    """
    accept_selectors = [
        (By.XPATH, "//button[.//span[contains(text(),'Aceitar tudo')]]"),
        (By.XPATH, "//button[contains(., 'Aceitar tudo')]"),
        (By.XPATH, "//button[contains(., 'Accept all')]"),
        (By.XPATH, "//button[contains(@aria-label, 'Aceitar')]"),
        (By.XPATH, "//button[contains(@aria-label, 'Accept')]"),
        (By.XPATH, "//form//button[2]"),  # fallback: geralmente o 2º botão do form de consentimento é "aceitar"
    ]

    short_wait = WebDriverWait(driver, 5)
    for by, selector in accept_selectors:
        try:
            button = short_wait.until(EC.element_to_be_clickable((by, selector)))
            _real_click(driver, button)
            if debug:
                print(f"[DEBUG] Tela de consentimento de cookies detectada e fechada usando: {selector}")
            time.sleep(1.5)
            return True
        except TimeoutException:
            continue
        except Exception as e:
            if debug:
                print(f"[DEBUG] Erro ao tentar clicar em '{selector}': {e}")
            continue

    if debug:
        print("[DEBUG] Nenhuma tela de consentimento de cookies detectada "
              "(ou ela já não apareceu dessa vez) — seguindo normalmente.")
    return False


def load_reviews_page(driver, url, timeout=20):
    """
    Abre a URL, fecha a tela de consentimento de cookies (se aparecer) e
    espera o painel de avaliações carregar.

    Retorna:
        True  -> página carregada e pronta para o scraping.
        False -> a página não carregou a tempo (um screenshot de erro é
                 salvo automaticamente para facilitar o diagnóstico).
    """
    driver.get(url)

    accept_cookie_consent(driver)

    try:
        WebDriverWait(driver, timeout).until(
            EC.presence_of_element_located((By.CLASS_NAME, 'jftiEf'))
        )
        return True
    except TimeoutException:
        error_base_path = os.path.dirname(os.path.abspath(__file__))
        screenshot_path = os.path.join(error_base_path, "erro_carregamento_avaliacoes.png")
        try:
            driver.save_screenshot(screenshot_path)
            print("ERRO: a página não carregou as avaliações dentro do tempo esperado.")
            print(f"Um screenshot da tela no momento da falha foi salvo em:\n  {screenshot_path}")
            print("Abra essa imagem para ver o que está bloqueando o carregamento "
                  "(ex: tela de consentimento, captcha, layout diferente do esperado).")
        except Exception as e:
            print(f"Não consegui nem salvar o screenshot de erro: {e}")
        return False


# ---------------------------------------------------------------------------
# Ordenação das avaliações
# ---------------------------------------------------------------------------

def sort_reviews_by_most_recent(driver, debug=True):
    """
    Ordena as avaliações do Google Maps por 'Mais recentes'.

    Retorna:
        True  -> ordenação realizada com sucesso.
        False -> não foi possível alterar a ordenação.
    """
    wait = WebDriverWait(driver, 15)

    try:
        wait.until(
            EC.presence_of_element_located(
                (By.CSS_SELECTOR, "div.m6QErb.DxyBCb.kA9KIf.dS8AEf")
            )
        )
    except TimeoutException:
        print("Painel de avaliações não carregou a tempo.")
        return False

    sort_button_selectors = [
        (By.XPATH, "//button[contains(@aria-label,'Classificar avaliações')]"),
        (By.XPATH, "//button[contains(@aria-label,'Ordenar')]"),
        (By.XPATH, "//button[.//span[contains(text(),'Ordenar')]]"),
        (By.XPATH, "//button[contains(.,'Ordenar')]"),
        (By.XPATH, "//button[contains(@aria-label,'Sort reviews')]"),
        (By.XPATH, "//button[.//span[contains(text(),'Sort')]]"),
        (By.XPATH, "//button[contains(.,'Sort')]"),
    ]

    sort_button = None
    for by, selector in sort_button_selectors:
        try:
            sort_button = wait.until(EC.element_to_be_clickable((by, selector)))
            if debug:
                print(f"[DEBUG] Botão de ordenação encontrado usando: {selector}")
            break
        except TimeoutException:
            continue

    if sort_button is None:
        print("Não foi possível localizar o botão de ordenação.")
        if debug:
            print("\n[DEBUG] Botões visíveis na página (para ajudar a achar o seletor certo):\n")
            for b in driver.find_elements(By.TAG_NAME, "button"):
                text = b.text.strip()
                aria = b.get_attribute("aria-label")
                if text or aria:
                    print(f"  texto='{text}' | aria-label='{aria}'")
        return False

    try:
        _real_click(driver, sort_button)
    except Exception as e:
        print(f"Erro ao clicar no botão de ordenação: {e}")
        return False

    try:
        wait.until(
            EC.presence_of_element_located((By.XPATH, "//div[@role='menu' or @role='listbox']"))
        )
        time.sleep(0.5)
    except TimeoutException:
        if debug:
            print("[DEBUG] O menu de ordenação não parece ter aberto (nenhum elemento com "
                  "role='menu'/'listbox' apareceu). O clique pode não ter surtido efeito.")

    most_recent_option_selectors = [
        (By.XPATH, "//div[@role='menuitemradio'][contains(.,'Mais recentes')]"),
        (By.XPATH, "//li[@role='menuitemradio' or @role='menuitem'][contains(.,'Mais recentes')]"),
        (By.XPATH, "//div[contains(.,'Mais recentes')]"),
        (By.XPATH, "//span[contains(.,'Mais recentes')]"),
        (By.XPATH, "//div[@role='menuitemradio'][contains(.,'Newest')]"),
        (By.XPATH, "//span[contains(.,'Newest')]"),
    ]

    most_recent_option = None
    for by, selector in most_recent_option_selectors:
        try:
            most_recent_option = wait.until(EC.element_to_be_clickable((by, selector)))
            if debug:
                print(f"[DEBUG] Opção 'Mais recentes' encontrada usando: {selector}")
            break
        except TimeoutException:
            continue

    if most_recent_option is None:
        if debug:
            menu_items = driver.find_elements(By.XPATH, "//*[@role='menuitemradio' or @role='menuitem']")
            print(f"[DEBUG] Não encontrei a opção 'Mais recentes'. Itens de menu visíveis ({len(menu_items)}):")
            for item in menu_items:
                print(f"  - texto: {item.text!r}")
        print("Não foi possível encontrar a opção 'Mais recentes'.")
        return False

    try:
        _real_click(driver, most_recent_option)
        time.sleep(2)
    except Exception as e:
        print(f"Erro ao clicar na opção 'Mais recentes': {e}")
        return False

    # Verificação extra: confirma que a ordenação realmente mudou, conferindo
    # se o texto do botão de ordenação mudou de "Mais relevantes" para
    # "Mais recentes" (evita confiar num clique que pode não ter surtido efeito)
    try:
        time.sleep(1)
        updated_button_text = sort_button.text.strip() if sort_button.text else (sort_button.get_attribute("aria-label") or "")
        if debug:
            print(f"[DEBUG] Texto do botão de ordenação após o clique: '{updated_button_text}'")
        if "recentes" in updated_button_text.lower() or "recent" in updated_button_text.lower():
            print("Confirmado: avaliações ordenadas por 'Mais recentes'.")
        else:
            print("Aviso: cliquei na opção, mas não consegui confirmar pelo texto do botão. "
                  "Verifique visualmente se a ordenação mudou.")
    except Exception:
        pass

    return True


# ---------------------------------------------------------------------------
# Distribuição de votos por estrela
# ---------------------------------------------------------------------------

def capture_star_distribution(driver, debug=True):
    """
    Captura, no cabeçalho do painel de avaliações, o número de votos para
    cada nota (5, 4, 3, 2 e 1 estrelas).

    Retorna um dicionário no formato:
        {5: 120, 4: 30, 3: 10, 2: 5, 1: 8}

    Se `debug=True`, imprime no console todos os aria-labels candidatos
    encontrados na página — útil se a extração continuar zerada, pois mostra
    exatamente o texto que precisa ser interpretado.
    """
    distribution = {5: 0, 4: 0, 3: 0, 2: 0, 1: 0}

    try:
        histogram_selectors = [
            "//div[contains(@aria-label, 'estrela')]",
            "//button[contains(@aria-label, 'estrela')]",
            "//*[contains(@aria-label, 'star')]",  # fallback em inglês
            "//div[contains(@aria-label,'5') and contains(@aria-label,'estrela')]",
        ]

        candidate_elements = []
        for xpath in histogram_selectors:
            candidate_elements.extend(driver.find_elements(By.XPATH, xpath))

        # remove duplicados mantendo a ordem
        seen = set()
        unique_elements = []
        for el in candidate_elements:
            if el.id not in seen:
                seen.add(el.id)
                unique_elements.append(el)

        if debug:
            print(f"[DEBUG] {len(unique_elements)} elementos candidatos ao histograma de estrelas encontrados.")

        for row in unique_elements:
            aria_label = row.get_attribute('aria-label')
            if not aria_label:
                continue

            if debug:
                print(f"[DEBUG] aria-label candidato: {aria_label!r}")

            aria_label_lower = aria_label.lower()

            found_star = None
            for n in range(5, 0, -1):
                if f"{n} estrela" in aria_label_lower or f"{n} star" in aria_label_lower:
                    found_star = n
                    break

            if found_star is None:
                continue

            numbers = [int(s) for s in aria_label_lower.replace(',', '').replace('.', '').split() if s.isdigit()]
            if len(numbers) >= 2:
                distribution[found_star] = numbers[1]
            elif len(numbers) == 1 and numbers[0] != found_star:
                distribution[found_star] = numbers[0]

        if all(v == 0 for v in distribution.values()):
            print("Não foi possível localizar a distribuição de estrelas na página "
                  "(veja os aria-labels de [DEBUG] acima para ajustar o seletor).")
        else:
            print(f"Distribuição de estrelas capturada: {distribution}")

    except Exception as e:
        print(f"Erro ao capturar distribuição de estrelas: {e}")

    return distribution


# ---------------------------------------------------------------------------
# Carregamento (scroll) e extração das avaliações
# ---------------------------------------------------------------------------

def load_reviews(driver):
    """
    Rola a lista de avaliações até o fim (sem limite fixo) e retorna a lista
    de elementos <div class="jftiEf"> (uma avaliação "crua", ainda como
    WebElement do Selenium — a extração dos dados de cada uma é feita
    separadamente por `extract_review_data`).

    A rolagem continua até que o "scrollHeight" da div pare de aumentar por
    várias tentativas seguidas, ou seja, até o Google Maps não ter mais
    nenhuma avaliação nova para carregar (lazy loading).
    """
    reviews = []
    seen_texts = set()

    try:
        scrollable_div = driver.find_element(By.CSS_SELECTOR, 'div.m6QErb.DxyBCb.kA9KIf.dS8AEf')
        last_height = driver.execute_script("return arguments[0].scrollHeight;", scrollable_div)

        max_attempts_without_change = 3
        attempts_without_change = 0

        while True:
            driver.execute_script("arguments[0].scrollTop = arguments[0].scrollHeight;", scrollable_div)
            time.sleep(2)
            new_height = driver.execute_script("return arguments[0].scrollHeight;", scrollable_div)

            if new_height == last_height:
                attempts_without_change += 1
                if attempts_without_change >= max_attempts_without_change:
                    print(f"Fim da lista de avaliações atingido. "
                          f"Total coletado: {len(reviews)} avaliações.")
                    break
            else:
                attempts_without_change = 0
                last_height = new_height

            new_reviews = driver.find_elements(By.CLASS_NAME, 'jftiEf')
            for review in new_reviews:
                review_text = review.text.strip()
                if review_text and review_text not in seen_texts:
                    reviews.append(review)
                    seen_texts.add(review_text)

            if len(reviews) % 50 == 0 and len(reviews) > 0:
                print(f"{len(reviews)} avaliações coletadas até agora...")

    except Exception as e:
        print(f"Erro ao carregar avaliações: {e}")

    return reviews


def extract_review_data(review):
    """
    Recebe um WebElement de uma avaliação (<div class="jftiEf">) e retorna
    um dicionário com os dados extraídos: nome, nota, data, comentário e
    resposta da empresa.

    Separar essa extração da geração de PDF/Excel deixa cada função com uma
    única responsabilidade: essa aqui só lê o DOM, as outras só escrevem
    arquivos.
    """
    try:
        try:
            more_button = review.find_element(By.CLASS_NAME, 'w8nwRe.kyuRq')
            more_button.click()
            time.sleep(1)
        except Exception:
            pass
        name = review.find_element(By.CLASS_NAME, 'd4r55').text.strip()
    except Exception:
        name = 'Nome não encontrado'

    try:
        rating = review.find_element(By.CLASS_NAME, 'kvMYJc').get_attribute('aria-label')
    except Exception:
        rating = 'Nota não encontrada'

    try:
        comment = review.find_element(By.CLASS_NAME, 'MyEned').text.strip()
    except Exception:
        comment = 'Comentário não encontrado'

    try:
        review_date = review.find_element(By.CLASS_NAME, 'rsqaWe').text.strip()
    except Exception:
        review_date = 'Data não encontrada'

    try:
        company_reply = review.find_element(By.XPATH, './/div[contains(@class, "wiI7pd")]').text.strip()
    except Exception:
        company_reply = 'Resposta não encontrada'

    return {
        'nome': name,
        'nota': rating,
        'data': review_date,
        'comentario': comment,
        'resposta_empresa': company_reply,
    }


# ---------------------------------------------------------------------------
# Geração dos arquivos de saída (PDF e Excel)
# ---------------------------------------------------------------------------

def generate_pdf(reviews_data, star_distribution, pdf_path):
    """Gera o arquivo PDF com todas as avaliações e a distribuição de estrelas."""
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Arial", size=12)

    for data in reviews_data:
        pdf.cell(200, 10, txt=safe_text(f"Nome: {data['nome']}"), ln=True, align='L')
        pdf.cell(200, 10, txt=safe_text(f"Nota: {data['nota']}"), ln=True, align='L')
        pdf.cell(200, 10, txt=safe_text(f"Data: {data['data']}"), ln=True, align='L')
        pdf.multi_cell(200, 10, txt=safe_text(f"Comentário: {data['comentario']}"), align='L')
        pdf.multi_cell(200, 10, txt=safe_text(f"Resposta da empresa: {data['resposta_empresa']}"), align='L')
        pdf.cell(200, 10, txt=safe_text('-' * 40), ln=True, align='L')

    # Página extra com a distribuição de avaliações por estrela
    pdf.add_page()
    pdf.set_font("Arial", size=14)
    pdf.cell(200, 10, txt=safe_text("Distribuição de avaliações por estrela"), ln=True, align='L')
    pdf.set_font("Arial", size=12)
    for star in sorted(star_distribution.keys(), reverse=True):
        pdf.cell(200, 10, txt=safe_text(f"{star} estrelas: {star_distribution[star]} avaliações"), ln=True, align='L')

    pdf.output(pdf_path)
    print(f"PDF salvo com sucesso em: {pdf_path}")


def generate_excel(reviews_data, star_distribution, excel_path):
    """Gera o arquivo Excel com uma aba de avaliações e outra de distribuição por estrela."""
    wb = Workbook()

    ws = wb.active
    ws.title = "Avaliações"
    ws.append(["Nome", "Nota", "Data", "Comentário", "Resposta da Empresa"])
    for data in reviews_data:
        ws.append([
            data['nome'],
            data['nota'],
            data['data'],
            data['comentario'],
            data['resposta_empresa'],
        ])

    ws_distribution = wb.create_sheet(title="Distribuição por Estrela")
    ws_distribution.append(["Estrelas", "Quantidade de Avaliações"])
    for star in sorted(star_distribution.keys(), reverse=True):
        ws_distribution.append([star, star_distribution[star]])

    wb.save(excel_path)
    print(f"Excel salvo com sucesso em: {excel_path}")


# ---------------------------------------------------------------------------
# Orquestração principal
# ---------------------------------------------------------------------------

def main():
    start_time = time.time()

    driver = configure_driver(headless=False)

    try:
        page_loaded = load_reviews_page(driver, REVIEWS_URL)
        if not page_loaded:
            return  # o erro e o screenshot já foram reportados dentro da função

        star_distribution = capture_star_distribution(driver)

        sort_success = sort_reviews_by_most_recent(driver)
        if not sort_success:
            print("ATENÇÃO: a ordenação por 'Mais recentes' falhou. "
                  "As avaliações serão coletadas na ordem padrão do Google (Mais relevantes).")

        review_elements = load_reviews(driver)
        reviews_data = [extract_review_data(review) for review in review_elements]

        # Pasta "avaliacoes" dentro do diretório do próprio programa
        base_path = os.path.dirname(os.path.abspath(__file__))
        documents_path = os.path.join(base_path, "avaliacoes")
        os.makedirs(documents_path, exist_ok=True)

        generate_pdf(reviews_data, star_distribution, os.path.join(documents_path, "avaliacoes.pdf"))
        generate_excel(reviews_data, star_distribution, os.path.join(documents_path, "avaliacoes.xlsx"))

    finally:
        # O navegador é fechado mesmo se algo der errado no meio do caminho
        driver.quit()

    end_time = time.time()
    minutes, seconds = divmod(end_time - start_time, 60)
    print(f"Tempo de execução do script: {int(minutes)} minutos e {seconds:.2f} segundos")


if __name__ == "__main__":
    main()