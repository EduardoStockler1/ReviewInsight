import os
import time
from fpdf import FPDF
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.action_chains import ActionChains
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import TimeoutException, NoSuchElementException
from webdriver_manager.chrome import ChromeDriverManager
from openpyxl import Workbook


# Limite de avaliações a coletar (None = sem limite, coleta tudo)
MAX_REVIEWS = 5

# URL da página de avaliações a raspar
REVIEWS_URL = (
    'https://www.google.com/maps/place/Hospital+Erastinho/'
    '@-25.452814,-49.2411939,1074m/data=!3m1!1e3!4m8!3m7!1s0x94dce5124cb3e99b:'
    '0xdb43625e49c502d9!8m2!3d-25.4528189!4d-49.238619!9m1!1b1!16s%2Fg%2F11h5mp07v2'
    '?entry=ttu&g_ep=EgoyMDI2MDYyOC4wIKXMDSoASAFQAw%3D%3D'
)


# =============================================================================
# 1) UTILITÁRIOS GERAIS (usados por várias etapas abaixo)
# =============================================================================

def safe_text(text):
    """Remove caracteres que o FPDF não sabe codificar."""
    return text.encode('latin-1', 'replace').decode('latin-1')


def _real_click(driver, element):
    """
    Clique "de verdade" (rola até o elemento + simula mouse com ActionChains).

    Necessário porque o Google Maps escuta mousedown/pointerdown, não só
    "click" — um clique via JS puro passa em branco sem dar erro nenhum.
    """
    driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", element)
    time.sleep(0.3)
    ActionChains(driver).move_to_element(element).pause(0.2).click().perform()


def get_available_directory(base_path, base_name):
    """
    Retorna um caminho de pasta livre: base_name, ou base_name1,
    base_name2... se já existir. Evita sobrescrever/travar em cima de uma
    pasta com arquivo já aberto de uma execução anterior.
    """
    dir_path = os.path.join(base_path, base_name)
    if not os.path.exists(dir_path):
        return dir_path

    counter = 1
    while True:
        new_dir_path = os.path.join(base_path, f"{base_name}{counter}")
        if not os.path.exists(new_dir_path):
            return new_dir_path
        counter += 1


# =============================================================================
# 2) ABRIR O NAVEGADOR
# =============================================================================

def configure_driver(headless=False):
    """Cria o Chrome WebDriver."""
    options = Options()
    if headless:
        options.add_argument('--headless')
    return webdriver.Chrome(service=Service(ChromeDriverManager().install()), options=options)


# =============================================================================
# 3) FECHAR TELA DE COOKIES (chamada de dentro de load_reviews_page)
# =============================================================================

def accept_cookie_consent(driver, debug=True):
    """Fecha o aviso de cookies do Google, se ele aparecer. Se não aparecer, segue normal."""
    accept_selectors = [
        (By.XPATH, "//button[.//span[contains(text(),'Aceitar tudo')]]"),
        (By.XPATH, "//button[contains(., 'Aceitar tudo')]"),
        (By.XPATH, "//button[contains(., 'Accept all')]"),
        (By.XPATH, "//button[contains(@aria-label, 'Aceitar')]"),
        (By.XPATH, "//button[contains(@aria-label, 'Accept')]"),
        (By.XPATH, "//form//button[2]"),  # fallback: 2º botão do form costuma ser "aceitar"
    ]

    short_wait = WebDriverWait(driver, 5)
    for by, selector in accept_selectors:
        try:
            button = short_wait.until(EC.element_to_be_clickable((by, selector)))
            _real_click(driver, button)
            if debug:
                print(f"[DEBUG] Cookies: fechado usando {selector}")
            time.sleep(1.5)
            return True
        except TimeoutException:
            continue
        except Exception as e:
            if debug:
                print(f"[DEBUG] Cookies: erro em '{selector}': {e}")
            continue

    if debug:
        print("[DEBUG] Cookies: nenhuma tela de consentimento apareceu.")
    return False


# =============================================================================
# 4) ABRIR A PÁGINA E ESPERAR AS AVALIAÇÕES CARREGAREM
# =============================================================================

def load_reviews_page(driver, url, timeout=20):
    """
    Abre a URL, fecha cookies e espera o painel de avaliações aparecer.
    Retorna True se carregou a tempo; False (+ screenshot de erro) se não.
    """
    driver.get(url)
    accept_cookie_consent(driver)

    try:
        WebDriverWait(driver, timeout).until(
            EC.presence_of_element_located((By.CLASS_NAME, 'jftiEf'))
        )
        return True
    except TimeoutException:
        error_base_path = os.path.dirname(os.path.abspath(__file__))
        screenshot_path = os.path.join(error_base_path, "erro_carregamento_avaliacoes.png")
        try:
            driver.save_screenshot(screenshot_path)
            print("ERRO: avaliações não carregaram a tempo.")
            print(f"Screenshot do erro salvo em: {screenshot_path}")
        except Exception as e:
            print(f"Não consegui nem salvar o screenshot: {e}")
        return False


# =============================================================================
# 5) LER A DISTRIBUIÇÃO DE ESTRELAS (5★, 4★, ... no topo do painel)
# =============================================================================

def capture_star_distribution(driver, debug=True):
    """Lê o histograma de estrelas do topo do painel. Retorna {5: n, 4: n, ...}."""
    distribution = {5: 0, 4: 0, 3: 0, 2: 0, 1: 0}

    try:
        histogram_selectors = [
            "//div[contains(@aria-label, 'estrela')]",
            "//button[contains(@aria-label, 'estrela')]",
            "//*[contains(@aria-label, 'star')]",  # fallback em inglês
        ]

        candidate_elements = []
        for xpath in histogram_selectors:
            candidate_elements.extend(driver.find_elements(By.XPATH, xpath))

        # remove duplicados mantendo a ordem
        seen = set()
        unique_elements = []
        for el in candidate_elements:
            if el.id not in seen:
                seen.add(el.id)
                unique_elements.append(el)

        if debug:
            print(f"[DEBUG] Estrelas: {len(unique_elements)} elementos candidatos encontrados.")

        for row in unique_elements:
            aria_label = row.get_attribute('aria-label')
            if not aria_label:
                continue
            if debug:
                print(f"[DEBUG] Estrelas: aria-label {aria_label!r}")

            aria_label_lower = aria_label.lower()
            found_star = None
            for n in range(5, 0, -1):
                if f"{n} estrela" in aria_label_lower or f"{n} star" in aria_label_lower:
                    found_star = n
                    break
            if found_star is None:
                continue

            numbers = [int(s) for s in aria_label_lower.replace(',', '').replace('.', '').split() if s.isdigit()]
            if len(numbers) >= 2:
                distribution[found_star] = numbers[1]
            elif len(numbers) == 1 and numbers[0] != found_star:
                distribution[found_star] = numbers[0]

        if all(v == 0 for v in distribution.values()):
            print("Distribuição de estrelas: não encontrada (veja os [DEBUG] acima).")
        else:
            print(f"Distribuição de estrelas capturada: {distribution}")

    except Exception as e:
        print(f"Erro ao capturar distribuição de estrelas: {e}")

    return distribution


# =============================================================================
# 6) ORDENAR AS AVALIAÇÕES POR "MAIS RECENTES"
# =============================================================================

def sort_reviews_by_most_recent(driver, debug=True):
    """Clica no botão de ordenação e seleciona 'Mais recentes'. Retorna True/False."""
    wait = WebDriverWait(driver, 15)

    try:
        wait.until(
            EC.presence_of_element_located((By.CSS_SELECTOR, "div.m6QErb.DxyBCb.kA9KIf.dS8AEf"))
        )
    except TimeoutException:
        print("Painel de avaliações não carregou a tempo.")
        return False

    sort_button_selectors = [
        (By.XPATH, "//button[contains(@aria-label,'Classificar avaliações')]"),
        (By.XPATH, "//button[contains(@aria-label,'Ordenar')]"),
        (By.XPATH, "//button[.//span[contains(text(),'Ordenar')]]"),
        (By.XPATH, "//button[contains(.,'Ordenar')]"),
        (By.XPATH, "//button[contains(@aria-label,'Sort reviews')]"),
        (By.XPATH, "//button[.//span[contains(text(),'Sort')]]"),
        (By.XPATH, "//button[contains(.,'Sort')]"),
    ]

    sort_button = None
    for by, selector in sort_button_selectors:
        try:
            sort_button = wait.until(EC.element_to_be_clickable((by, selector)))
            if debug:
                print(f"[DEBUG] Ordenação: botão encontrado com {selector}")
            break
        except TimeoutException:
            continue

    if sort_button is None:
        print("Não achei o botão de ordenação.")
        return False

    try:
        _real_click(driver, sort_button)
    except Exception as e:
        print(f"Erro ao clicar no botão de ordenação: {e}")
        return False

    try:
        wait.until(EC.presence_of_element_located((By.XPATH, "//div[@role='menu' or @role='listbox']")))
        time.sleep(0.5)
    except TimeoutException:
        if debug:
            print("[DEBUG] Ordenação: o menu não parece ter aberto.")

    most_recent_option_selectors = [
        (By.XPATH, "//div[@role='menuitemradio'][contains(.,'Mais recentes')]"),
        (By.XPATH, "//li[@role='menuitemradio' or @role='menuitem'][contains(.,'Mais recentes')]"),
        (By.XPATH, "//div[contains(.,'Mais recentes')]"),
        (By.XPATH, "//span[contains(.,'Mais recentes')]"),
        (By.XPATH, "//div[@role='menuitemradio'][contains(.,'Newest')]"),
        (By.XPATH, "//span[contains(.,'Newest')]"),
    ]

    most_recent_option = None
    for by, selector in most_recent_option_selectors:
        try:
            most_recent_option = wait.until(EC.element_to_be_clickable((by, selector)))
            if debug:
                print(f"[DEBUG] Ordenação: opção 'Mais recentes' encontrada com {selector}")
            break
        except TimeoutException:
            continue

    if most_recent_option is None:
        print("Não achei a opção 'Mais recentes'.")
        return False

    try:
        _real_click(driver, most_recent_option)
        time.sleep(2)
    except Exception as e:
        print(f"Erro ao clicar em 'Mais recentes': {e}")
        return False

    # confere se o texto do botão realmente mudou (evita clique "fantasma")
    try:
        time.sleep(1)
        updated_text = sort_button.text.strip() if sort_button.text else (sort_button.get_attribute("aria-label") or "")
        if "recentes" in updated_text.lower() or "recent" in updated_text.lower():
            print("Confirmado: avaliações ordenadas por 'Mais recentes'.")
        else:
            print("Aviso: cliquei, mas não confirmei pelo texto do botão.")
    except Exception:
        pass

    return True


# =============================================================================
# 7) ROLAR A LISTA E CONTAR QUANTAS AVALIAÇÕES CARREGARAM
# =============================================================================

def load_reviews(driver, max_reviews=None):
    """
    Rola até o fim (ou até max_reviews) e retorna a CONTAGEM de avaliações
    carregadas — não os elementos em si, pois eles ficam "obsoletos" quando
    clicamos no "Mais" depois (ver extract_review_data).
    """
    try:
        scrollable_div = driver.find_element(By.CSS_SELECTOR, 'div.m6QErb.DxyBCb.kA9KIf.dS8AEf')
        last_height = driver.execute_script("return arguments[0].scrollHeight;", scrollable_div)

        max_attempts_without_change = 3
        attempts_without_change = 0

        while True:
            if max_reviews is not None:
                current_count = len(driver.find_elements(By.CLASS_NAME, 'jftiEf'))
                if current_count >= max_reviews:
                    print(f"Limite de {max_reviews} avaliações atingido.")
                    break

            driver.execute_script("arguments[0].scrollTop = arguments[0].scrollHeight;", scrollable_div)
            time.sleep(2)
            new_height = driver.execute_script("return arguments[0].scrollHeight;", scrollable_div)

            if new_height == last_height:
                attempts_without_change += 1
                if attempts_without_change >= max_attempts_without_change:
                    break
            else:
                attempts_without_change = 0
                last_height = new_height

            current_count = len(driver.find_elements(By.CLASS_NAME, 'jftiEf'))
            if current_count % 50 == 0 and current_count > 0:
                print(f"{current_count} avaliações carregadas até agora...")

    except Exception as e:
        print(f"Erro ao rolar avaliações: {e}")

    total = len(driver.find_elements(By.CLASS_NAME, 'jftiEf'))
    if max_reviews is not None:
        total = min(total, max_reviews)

    print(f"Fim da coleta. Total: {total} avaliações.")
    return total


# =============================================================================
# 8) EXTRAIR OS DADOS DE CADA AVALIAÇÃO (chamada uma vez por índice)
# =============================================================================

def extract_review_data(driver, index):
    """
    Extrai nome, nota, data, comentário e resposta da empresa da avaliação
    na posição `index`. Antes disso, expande todo texto truncado clicando
    em qualquer botão "Mais" que exista.
    """

    def get_review_element():
        """Re-busca o elemento pelo índice (nunca reusa referência antiga)."""
        elements = driver.find_elements(By.CLASS_NAME, 'jftiEf')
        return elements[index] if index < len(elements) else None

    default_result = {
        'nome': 'Nome não encontrado',
        'nota': 'Nota não encontrada',
        'data': 'Data não encontrada',
        'comentario': 'Comentário não encontrado',
        'resposta_empresa': 'Resposta não encontrada',
    }

    # ---- Expande todos os "Mais" (comentário e/ou resposta da empresa) ----
    #
    # Não dá pra calcular "quantos botões existem" uma vez só: depois de
    # clicado, o botão MUDA DE CLASSE e some do seletor. Por isso sempre
    # pegamos só o PRIMEIRO botão restante e clicamos, repetindo até não
    # sobrar nenhum — assim funciona com 0, 1 ou 2 botões, e não depende
    # de índice fixo numa lista que encolhe a cada clique.
    max_click_attempts = 5  # trava de segurança contra loop infinito
    for _ in range(max_click_attempts):
        review = get_review_element()
        if review is None:
            break
        remaining_buttons = review.find_elements(By.CSS_SELECTOR, '.w8nwRe.kyuRq')
        if not remaining_buttons:
            break
        try:
            _real_click(driver, remaining_buttons[0])
        except Exception:
            break

    # ---- Re-busca a avaliação (o clique pode ter renderizado ela de novo) ----
    review = get_review_element()
    if review is None:
        return default_result

    try:
        name = review.find_element(By.CLASS_NAME, 'd4r55').text.strip()
    except Exception:
        name = default_result['nome']

    try:
        rating = review.find_element(By.CLASS_NAME, 'kvMYJc').get_attribute('aria-label')
    except Exception:
        rating = default_result['nota']

    try:
        comment = review.find_element(By.CLASS_NAME, 'MyEned').text.strip()
    except Exception:
        comment = default_result['comentario']

    try:
        review_date = review.find_element(By.CLASS_NAME, 'rsqaWe').text.strip()
    except Exception:
        review_date = default_result['data']

    try:
        company_reply = review.find_element(By.XPATH, './/div[contains(@class, "wiI7pd")]').text.strip()
    except Exception:
        company_reply = default_result['resposta_empresa']

    return {
        'nome': name,
        'nota': rating,
        'data': review_date,
        'comentario': comment,
        'resposta_empresa': company_reply,
    }


# =============================================================================
# 9) GERAR O PDF
# =============================================================================

def generate_pdf(reviews_data, star_distribution, pdf_path):
    """Escreve todas as avaliações + a distribuição de estrelas num PDF."""
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Arial", size=12)

    for data in reviews_data:
        pdf.cell(200, 10, txt=safe_text(f"Nome: {data['nome']}"), ln=True, align='L')
        pdf.cell(200, 10, txt=safe_text(f"Nota: {data['nota']}"), ln=True, align='L')
        pdf.cell(200, 10, txt=safe_text(f"Data: {data['data']}"), ln=True, align='L')
        pdf.multi_cell(200, 10, txt=safe_text(f"Comentário: {data['comentario']}"), align='L')
        pdf.multi_cell(200, 10, txt=safe_text(f"Resposta da empresa: {data['resposta_empresa']}"), align='L')
        pdf.cell(200, 10, txt=safe_text('-' * 40), ln=True, align='L')

    pdf.add_page()
    pdf.set_font("Arial", size=14)
    pdf.cell(200, 10, txt=safe_text("Distribuição de avaliações por estrela"), ln=True, align='L')
    pdf.set_font("Arial", size=12)
    for star in sorted(star_distribution.keys(), reverse=True):
        pdf.cell(200, 10, txt=safe_text(f"{star} estrelas: {star_distribution[star]} avaliações"), ln=True, align='L')

    pdf.output(pdf_path)
    print(f"PDF salvo com sucesso em: {pdf_path}")


# =============================================================================
# 10) GERAR O EXCEL
# =============================================================================

def generate_excel(reviews_data, star_distribution, excel_path):
    """Escreve uma aba de avaliações + uma aba de distribuição por estrela num Excel."""
    wb = Workbook()

    ws = wb.active
    ws.title = "Avaliações"
    ws.append(["Nome", "Nota", "Data", "Comentário", "Resposta da Empresa"])
    for data in reviews_data:
        ws.append([data['nome'], data['nota'], data['data'], data['comentario'], data['resposta_empresa']])

    ws_distribution = wb.create_sheet(title="Distribuição por Estrela")
    ws_distribution.append(["Estrelas", "Quantidade de Avaliações"])
    for star in sorted(star_distribution.keys(), reverse=True):
        ws_distribution.append([star, star_distribution[star]])

    wb.save(excel_path)
    print(f"Excel salvo com sucesso em: {excel_path}")


# =============================================================================
# 11) ORQUESTRAÇÃO — chama tudo acima, na ordem
# =============================================================================

def main():
    start_time = time.time()

    driver = configure_driver(headless=False)

    try:
        if not load_reviews_page(driver, REVIEWS_URL):
            return  # erro e screenshot já reportados dentro da função

        star_distribution = capture_star_distribution(driver)

        if not sort_reviews_by_most_recent(driver):
            print("ATENÇÃO: ordenação falhou — coletando na ordem padrão (Mais relevantes).")

        review_count = load_reviews(driver, max_reviews=MAX_REVIEWS)
        reviews_data = [extract_review_data(driver, i) for i in range(review_count)]

        # pasta "avaliacoes" (ou avaliacoes1, avaliacoes2... se já existir)
        base_path = os.path.dirname(os.path.abspath(__file__))
        documents_path = get_available_directory(base_path, "avaliacoes")
        os.makedirs(documents_path, exist_ok=True)

        generate_pdf(reviews_data, star_distribution, os.path.join(documents_path, "avaliacoes.pdf"))
        generate_excel(reviews_data, star_distribution, os.path.join(documents_path, "avaliacoes.xlsx"))

    finally:
        driver.quit()

    minutes, seconds = divmod(time.time() - start_time, 60)
    print(f"Tempo de execução: {int(minutes)} min e {seconds:.2f} s")


if __name__ == "__main__":
    main()