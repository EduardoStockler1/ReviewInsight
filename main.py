import os
import csv
import scraper
from fpdf import FPDF
import time
from openpyxl import Workbook
from config import MAX_REVIEWS, REVIEWS_URL

def get_available_directory(base_path, base_name):
    # Retorna um caminho de pasta livre: base_name, ou base_name1,
    # base_name2... se já existir. Evita sobrescrever/travar em cima de uma
    # pasta com arquivo já aberto de uma execução anterior.
    
    dir_path = os.path.join(base_path, base_name)
    if not os.path.exists(dir_path):
        return dir_path

    counter = 1
    while True:
        new_dir_path = os.path.join(base_path, f"{base_name}{counter}")
        if not os.path.exists(new_dir_path):
            return new_dir_path
        counter += 1



def generate_pdf(reviews_data, star_distribution, pdf_path):
    # Escreve todas as avaliações + a distribuição de estrelas num PDF.
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Arial", size=12)

    for data in reviews_data:
        pdf.cell(200, 10, txt=scraper.safe_text(f"Nome: {data['nome']}"), ln=True, align='L')
        pdf.cell(200, 10, txt=scraper.safe_text(f"Nota: {data['nota']}"), ln=True, align='L')
        pdf.cell(200, 10, txt=scraper.safe_text(f"Data: {data['data']}"), ln=True, align='L')
        pdf.multi_cell(200, 10, txt=scraper.safe_text(f"Comentário: {data['comentario']}"), align='L')
        pdf.multi_cell(200, 10, txt=scraper.safe_text(f"Resposta da empresa: {data['resposta_empresa']}"), align='L')
        pdf.cell(200, 10, txt=scraper.safe_text('-' * 40), ln=True, align='L')

    pdf.add_page()
    pdf.set_font("Arial", size=14)
    pdf.cell(200, 10, txt=scraper.safe_text("Distribuição de avaliações por estrela"), ln=True, align='L')
    pdf.set_font("Arial", size=12)
    for star in sorted(star_distribution.keys(), reverse=True):
        pdf.cell(200, 10, txt=scraper.safe_text(f"{star} estrelas: {star_distribution[star]} avaliações"), ln=True, align='L')

    pdf.output(pdf_path)
    print(f"PDF salvo com sucesso em: {pdf_path}")


def generate_excel(reviews_data, star_distribution, excel_path):
    # Escreve uma aba de avaliações + uma aba de distribuição por estrela num Excel.
    wb = Workbook()

    ws = wb.active
    ws.title = "Avaliações"
    ws.append(["Nome", "Nota", "Data", "Comentário", "Resposta da Empresa"])
    for data in reviews_data:
        ws.append([data['nome'], data['nota'], data['data'], data['comentario'], data['resposta_empresa']])

    ws_distribution = wb.create_sheet(title="Distribuição por Estrela")
    ws_distribution.append(["Estrelas", "Quantidade de Avaliações"])
    for star in sorted(star_distribution.keys(), reverse=True):
        ws_distribution.append([star, star_distribution[star]])

    wb.save(excel_path)
    print(f"Excel salvo com sucesso em: {excel_path}")

def generate_csv(reviews_data, csv_path):
    # Escreve as avaliações num CSV, uma linha por avaliação.
    with open(csv_path, mode='w', newline='', encoding='utf-8-sig') as f:
        writer = csv.writer(f)
        writer.writerow(["Nome", "Nota", "Data", "Comentário", "Resposta da Empresa"])
        for data in reviews_data:
            writer.writerow([
                data['nome'],
                data['nota'],
                data['data'],
                data['comentario'],
                data['resposta_empresa'],
            ])

    print(f"CSV salvo com sucesso em: {csv_path}")

def generate_star_distribution_csv(star_distribution, csv_path):
    with open(csv_path, mode='w', newline='', encoding='utf-8-sig') as f:
        writer = csv.writer(f)
        writer.writerow(["Estrelas", "Quantidade de Avaliações"])
        for star in sorted(star_distribution.keys(), reverse=True):
            writer.writerow([star, star_distribution[star]])

    print(f"CSV de distribuição salvo com sucesso em: {csv_path}")

def main():
    start_time = time.time()

    driver = scraper.configure_driver(headless=False)

    try:
        if not scraper.load_reviews_page(driver, REVIEWS_URL):
            return  # erro e screenshot já reportados dentro da função
            print("Abacaxi")

        star_distribution = scraper.capture_star_distribution(driver)

        if not scraper.sort_reviews_by_most_recent(driver):
            print("ATENÇÃO: ordenação falhou — coletando na ordem padrão (Mais relevantes).")


        review_count = scraper.load_reviews(driver, max_reviews=MAX_REVIEWS)
        reviews_data = [scraper.extract_review_data(driver, i) for i in range(review_count)]

        # pasta "avaliacoes" (ou avaliacoes1, avaliacoes2... se já existir)
        base_path = os.path.dirname(os.path.abspath(__file__))
        avaliacoes_path = os.path.join(base_path, "avaliacoes")
        documents_path = get_available_directory(avaliacoes_path, "avaliacoes")
        os.makedirs(documents_path, exist_ok=True)

        # generate_pdf(reviews_data, star_distribution, os.path.join(documents_path, "avaliacoes.pdf"))
        # generate_excel(reviews_data, star_distribution, os.path.join(documents_path, "avaliacoes.xlsx"))
        generate_csv(reviews_data, os.path.join(documents_path, "avaliacoes.csv"))
        generate_star_distribution_csv(star_distribution, os.path.join(documents_path, "distribuicao_estrelas.csv"))
    finally:
        driver.quit()

    minutes, seconds = divmod(time.time() - start_time, 60)
    print(f"Tempo de execução: {int(minutes)} min e {seconds:.2f} s")

if __name__ == "__main__":
    main()